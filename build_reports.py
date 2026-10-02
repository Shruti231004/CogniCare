import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

from reportlab.lib.pagesizes import letter
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Image as RLImage, Table as RLTable, TableStyle, PageBreak, KeepTogether, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def build_docx_report():
    doc = docx.Document()

    # Margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base styling
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    # ---------------- PAGE 1: TITLE PAGE ----------------
    p_exp = doc.add_paragraph()
    p_exp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_exp.paragraph_format.space_before = Pt(0)
    p_exp.paragraph_format.space_after = Pt(12)
    run_exp = p_exp.add_run("EXP 10")
    run_exp.font.name = "Times New Roman"
    run_exp.font.size = Pt(14)
    run_exp.font.bold = True

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(12)
    run_sub = p_sub.add_run("Report of Mini-project On")
    run_sub.font.name = "Times New Roman"
    run_sub.font.size = Pt(14)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(18)
    run_title = p_title.add_run("CogniCare AI: Brain Stroke Risk Prediction, MRI Lesion Segmentation and Emergency Response Platform")
    run_title.font.name = "Times New Roman"
    run_title.font.size = Pt(16)
    run_title.font.bold = True

    p_partial = doc.add_paragraph()
    p_partial.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_partial.paragraph_format.space_after = Pt(4)
    run_partial1 = p_partial.add_run("Submitted in partial fulfillment of the requirements of the Mini project in the\n")
    run_partial1.font.name = "Times New Roman"
    run_partial1.font.size = Pt(12)
    run_partial2 = p_partial.add_run("Subject: AI FOR HEALTHCARE of\n")
    run_partial2.font.name = "Times New Roman"
    run_partial2.font.size = Pt(12)
    run_partial2.font.bold = True
    run_partial3 = p_partial.add_run("Semester VII, FINAL Year Artificial Intelligence and Data Science")
    run_partial3.font.name = "Times New Roman"
    run_partial3.font.size = Pt(12)

    p_by = doc.add_paragraph()
    p_by.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_by.paragraph_format.space_before = Pt(14)
    p_by.paragraph_format.space_after = Pt(4)
    run_by = p_by.add_run("by")
    run_by.font.name = "Times New Roman"
    run_by.font.size = Pt(12)

    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_name.paragraph_format.space_after = Pt(16)
    run_name = p_name.add_run("Shruti Gauchandra (Roll No. 15)")
    run_name.font.name = "Times New Roman"
    run_name.font.size = Pt(13)
    run_name.font.bold = True

    # University Logo
    if os.path.exists("page_1_img_1.png"):
        p_logo1 = doc.add_paragraph()
        p_logo1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo1.paragraph_format.space_after = Pt(8)
        p_logo1.add_run().add_picture("page_1_img_1.png", width=Inches(1.2))

    p_univ = doc.add_paragraph()
    p_univ.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_univ.paragraph_format.space_after = Pt(4)
    run_univ = p_univ.add_run("University of Mumbai")
    run_univ.font.name = "Times New Roman"
    run_univ.font.size = Pt(13)
    run_univ.font.bold = True

    p_coll = doc.add_paragraph()
    p_coll.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_coll.paragraph_format.space_after = Pt(2)
    run_coll = p_coll.add_run("Vidyavardhini's College of Engineering & Technology")
    run_coll.font.name = "Times New Roman"
    run_coll.font.size = Pt(12)
    run_coll.font.bold = True

    p_dept = doc.add_paragraph()
    p_dept.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_dept.paragraph_format.space_after = Pt(10)
    run_dept = p_dept.add_run("Department of Artificial Intelligence and Data Science")
    run_dept.font.name = "Times New Roman"
    run_dept.font.size = Pt(12)
    run_dept.font.bold = True

    # College Logo
    if os.path.exists("page_1_img_2.png"):
        p_logo2 = doc.add_paragraph()
        p_logo2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo2.paragraph_format.space_after = Pt(8)
        p_logo2.add_run().add_picture("page_1_img_2.png", width=Inches(1.1))

    p_ay = doc.add_paragraph()
    p_ay.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ay.paragraph_format.space_after = Pt(0)
    run_ay = p_ay.add_run("(A.Y. 2026-27)")
    run_ay.font.name = "Times New Roman"
    run_ay.font.size = Pt(12)
    run_ay.font.bold = True

    doc.add_page_break()

    # ---------------- PAGE 2: CERTIFICATE ----------------
    p_cert_title = doc.add_paragraph()
    p_cert_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cert_title.paragraph_format.space_before = Pt(36)
    p_cert_title.paragraph_format.space_after = Pt(28)
    run_cert_title = p_cert_title.add_run("CERTIFICATE")
    run_cert_title.font.name = "Times New Roman"
    run_cert_title.font.size = Pt(16)
    run_cert_title.font.bold = True

    p_cert_body = doc.add_paragraph()
    p_cert_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_cert_body.paragraph_format.line_spacing = 1.4
    p_cert_body.paragraph_format.space_after = Pt(50)

    run_cb1 = p_cert_body.add_run('This is to certify that the Mini Project entitled ')
    run_cb1.font.name = "Times New Roman"
    run_cb1.font.size = Pt(12)

    run_cb2 = p_cert_body.add_run('"CogniCare AI: Brain Stroke Risk Prediction, MRI Lesion Segmentation and Emergency Response Platform"')
    run_cb2.font.name = "Times New Roman"
    run_cb2.font.size = Pt(12)
    run_cb2.font.bold = True

    run_cb3 = p_cert_body.add_run(' is submitted by ')
    run_cb3.font.name = "Times New Roman"
    run_cb3.font.size = Pt(12)

    run_cb4 = p_cert_body.add_run("Shruti Gauchandra (Roll No. 15)")
    run_cb4.font.name = "Times New Roman"
    run_cb4.font.size = Pt(12)
    run_cb4.font.bold = True

    run_cb5 = p_cert_body.add_run(" for the subject of ")
    run_cb5.font.name = "Times New Roman"
    run_cb5.font.size = Pt(12)

    run_cb6 = p_cert_body.add_run("AI for Healthcare")
    run_cb6.font.name = "Times New Roman"
    run_cb6.font.size = Pt(12)
    run_cb6.font.bold = True

    run_cb7 = p_cert_body.add_run(" in the ")
    run_cb7.font.name = "Times New Roman"
    run_cb7.font.size = Pt(12)

    run_cb8 = p_cert_body.add_run("Department of Artificial Intelligence and Data Science")
    run_cb8.font.name = "Times New Roman"
    run_cb8.font.size = Pt(12)
    run_cb8.font.bold = True

    run_cb9 = p_cert_body.add_run(" as a record of work done by him/her under our supervision and guidance.")
    run_cb9.font.name = "Times New Roman"
    run_cb9.font.size = Pt(12)

    p_guide = doc.add_paragraph()
    p_guide.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_guide.paragraph_format.space_before = Pt(30)
    p_guide.paragraph_format.space_after = Pt(2)
    run_guide = p_guide.add_run("Guide\n")
    run_guide.font.name = "Times New Roman"
    run_guide.font.size = Pt(13)
    run_guide.font.bold = True

    run_guide_name = p_guide.add_run("Asst. Prof. Kranti Gule\n")
    run_guide_name.font.name = "Times New Roman"
    run_guide_name.font.size = Pt(12)
    run_guide_name.font.bold = True

    run_guide_dept = p_guide.add_run("Department of Artificial Intelligence and Data Science\nVCET, Vasai")
    run_guide_dept.font.name = "Times New Roman"
    run_guide_dept.font.size = Pt(11)

    doc.add_page_break()

    # ---------------- PAGE 3: CONTENTS / INDEX ----------------
    p_contents = doc.add_paragraph()
    p_contents.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_contents.paragraph_format.space_before = Pt(20)
    p_contents.paragraph_format.space_after = Pt(20)
    run_contents = p_contents.add_run("Contents")
    run_contents.font.name = "Times New Roman"
    run_contents.font.size = Pt(16)
    run_contents.font.bold = True

    table = doc.add_table(rows=5, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_widths = [Inches(1.0), Inches(4.5), Inches(1.0)]
    headers = ["Sr. No", "Title", "Page No"]
    rows_data = [
        ["1", "Introduction of Project", "4"],
        ["2", "Importance of Project", "5"],
        ["3", "Screen Shot of Output", "6"],
        ["4", "Conclusion", "7"]
    ]

    hdr_cells = table.rows[0].cells
    for i, h_text in enumerate(headers):
        hdr_cells[i].text = h_text
        hdr_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = hdr_cells[i].paragraphs[0].runs[0]
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run.font.bold = True
        set_cell_background(hdr_cells[i], "F1F5F9")

    for r_idx, row in enumerate(rows_data):
        cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            cells[c_idx].text = val
            p = cells[c_idx].paragraphs[0]
            if c_idx == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.runs[0]
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)
            if r_idx % 2 == 1:
                set_cell_background(cells[c_idx], "F8FAFC")

    for row in table.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = w

    doc.add_page_break()

    # ---------------- PAGE 4: INTRODUCTION ----------------
    p_sec1 = doc.add_paragraph()
    p_sec1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sec1.paragraph_format.space_before = Pt(8)
    p_sec1.paragraph_format.space_after = Pt(10)
    run_sec1 = p_sec1.add_run("INTRODUCTION")
    run_sec1.font.name = "Times New Roman"
    run_sec1.font.size = Pt(15)
    run_sec1.font.bold = True

    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p1.paragraph_format.space_after = Pt(5)
    p1.add_run(
        "Stroke is one of the leading global causes of mortality and acquired adult long-term disability, "
        "accounting for approximately 11% of all deaths annually according to the World Health Organization (WHO). "
        "An acute ischemic stroke (AIS) occurs when a thrombus or embolus occludes cerebral blood vessels, depriving downstream "
        "neural tissue of oxygen and glucose. In clinical stroke neurology, the widely recognized paradigm is 'Time is Brain'—for "
        "every minute an ischemic stroke remains untreated, an estimated 1.9 million neurons, 14 billion synapses, and 12 km "
        "of myelinated nerve fibers undergo irreversible necrosis. Despite the efficacy of contemporary reperfusion therapies "
        "such as intravenous tissue plasminogen activator (IV rtPA) and endovascular thrombectomy, systemic bottlenecks in symptom "
        "recognition, clinical risk estimation, radiologic interpretation, and emergency center navigation lead to severe delays "
        "and suboptimal patient outcomes."
    )

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p2.paragraph_format.space_after = Pt(5)
    p2.add_run(
        "To mitigate these critical challenges, CogniCare AI is developed as an end-to-end, multi-modal artificial intelligence "
        "platform dedicated to comprehensive brain stroke risk management, automated neuroimaging analysis, and rapid emergency dispatch. "
        "Designed in strict accordance with the AI for Healthcare curriculum, the system bridges the gap between predictive statistical modeling, "
        "computer vision, medical natural language processing (NLP), and frontline clinical workflows. "
        "CogniCare AI is architected into six key cohesive modules:"
    )

    modules_info = [
        ("1. Public Awareness & FAST Screening Engine: ", 
         "An interactive digital screening assessment implementing the clinical Face, Arm, Speech, and Time (FAST) diagnostic protocol. "
         "The tool captures self-reported or bystander-observed deficits and immediately categorizes the suspicion index, initiating priority emergency pathways when acute focal neurological deficits are present."),
        ("2. Soft-Voting Machine Learning Ensemble: ", 
         "A tripartite ensemble fusing Random Forest (n=100 estimators), Gradient Boosting (learning rate = 0.08), and AdaBoost classifiers. "
         "Trained on stratified clinical parameters (including age, hypertension, heart disease, average blood glucose level, BMI, work type, and smoking status), "
         "the ensemble computes well-calibrated stroke probability estimates with high sensitivity."),
        ("3. 5-Stage MRI Brain Scan Segmentation Pipeline: ", 
         "A deterministic computer vision workflow engineered for axial T2-FLAIR and DWI brain MRI scans. The pipeline sequentially applies: "
         "(i) Grayscale intensity transformation, (ii) Gaussian spatial smoothing (5x5 kernel, sigma=1.2) to attenuate acquisition noise, "
         "(iii) Otsu's optimal bimodal thresholding to isolate hyperintense ischemic territories, (iv) Mathematical morphological opening and closing "
         "to eliminate skull-bone artifacts and ventricular noise, and (v) Telemetric quantification computing segmented lesion area (pixels), "
         "lesion burden ratio (%), and automated ASPECTS anatomical territory localization."),
        ("4. Clinical Natural Language Processing (NLP) Assistant: ", 
         "A medical text parsing engine that processes unstructured doctor, triage, and nurse EHR progress notes. "
         "Leveraging regex tokenization, clinical synonym dictionaries, and fuzzy string matching, the module extracts symptoms, diagnoses, "
         "and pharmacological prescriptions, automatically mapping them to standardized ontological codes (ICD-10, RxNorm, and SNOMED-CT)."),
        ("5. Explainable AI (XAI) Biomarker Attribution: ", 
         "Provides transparent, SHAP-inspired local feature contribution scores for each prediction, highlighting critical clinical biomarker alerts "
         "(e.g., severe hyperglycemia >200 mg/dL, stage-2 hypertension) to foster clinical trust and physician-in-the-loop validation."),
        ("6. Stroke Emergency Dispatch & Telemetry: ", 
         "A one-tap emergency protocol integrating a 4.5-hour intravenous thrombolysis countdown timer, geolocation-based nearest stroke-certified "
         "comprehensive hospital navigation, and paramedic handoff cards."
        )
    ]

    for title_text, desc_text in modules_info:
        p_mod = doc.add_paragraph()
        p_mod.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_mod.paragraph_format.space_after = Pt(4)
        r_t = p_mod.add_run(title_text)
        r_t.font.bold = True
        r_t.font.size = Pt(11)
        r_d = p_mod.add_run(desc_text)
        r_d.font.size = Pt(11)

    doc.add_page_break()

    # ---------------- PAGE 5: IMPORTANCE OF PROJECT ----------------
    p_sec2 = doc.add_paragraph()
    p_sec2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sec2.paragraph_format.space_before = Pt(8)
    p_sec2.paragraph_format.space_after = Pt(10)
    run_sec2 = p_sec2.add_run("IMPORTANCE OF PROJECT")
    run_sec2.font.name = "Times New Roman"
    run_sec2.font.size = Pt(15)
    run_sec2.font.bold = True

    imp_points = [
        ("1. Critical Golden Hour Optimization and Door-to-Needle Reduction",
         "In acute ischemic stroke, therapeutic efficacy is governed by the therapeutic golden window: intravenous thrombolysis (IV rtPA) must "
         "be administered within 4.5 hours from the last known well time, and endovascular thrombectomy is indicated within 6 to 24 hours. "
         "Every 15-minute reduction in door-to-needle time results in a 5% relative reduction in 90-day mortality. CogniCare AI directly "
         "compresses this critical timeline through instantaneous FAST identification, immediate emergency dispatch, and live countdown telemetry."),

        ("2. Objective, Deterministic, and Quantitative Neuroimaging Assessment",
         "Manual radiologic interpretation of acute brain scans is prone to inter-observer variability, especially in primary care settings without "
         "24/7 on-call neuroradiologists. The 5-stage automated segmentation pipeline delivers standardized, objective metrics of ischemic core burden "
         "and territory localization within milliseconds, preventing diagnostic hesitation and supporting emergency triage decisions."),

        ("3. Eliminating Black-Box Hesitation via Explainable AI (XAI)",
         "A primary barrier to the clinical adoption of machine learning in healthcare is the opaque nature of complex models. "
         "Clinicians cannot ethically or legally act on unverified predictions. CogniCare AI addresses this challenge by providing granular feature "
         "attributions and transparent voting consensus across three distinct ensemble classifiers, demonstrating clearly whether risk is driven "
         "by vascular hypertension, diabetic microangiopathy, or advancing age."),

        ("4. Streamlining Clinical Workflows & EHR Interoperability",
         "Emergency physicians face severe documentation burden, spending up to two hours on administrative electronic health record (EHR) entries "
         "for every hour of direct patient care. CogniCare AI's clinical NLP assistant instantly normalizes free-text clinical notes into internationally "
         "recognized ICD-10 and SNOMED-CT taxonomies, accelerating patient handoffs between emergency departments and neuro-intensive care units."),

        ("5. Democratizing Advanced Stroke Care in Rural and Underserved Facilities",
         "Tertiary stroke centers are disproportionately clustered in major metropolitan hubs. Rural and community district hospitals often lack "
         "sub-specialty stroke teams. CogniCare AI acts as a sophisticated clinical force multiplier, equipping medical officers with specialist-level "
         "decision support and rapid transfer coordination protocols to ensure equitable healthcare delivery.")
    ]

    for h_txt, b_txt in imp_points:
        p_imp = doc.add_paragraph()
        p_imp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_imp.paragraph_format.space_after = Pt(6)
        r_h = p_imp.add_run(h_txt + "\n")
        r_h.font.bold = True
        r_h.font.size = Pt(11.5)
        r_b = p_imp.add_run(b_txt)
        r_b.font.size = Pt(11)

    doc.add_page_break()

    # ---------------- PAGE 6: SCREEN SHOT OF OUTPUT ----------------
    p_sec3 = doc.add_paragraph()
    p_sec3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sec3.paragraph_format.space_before = Pt(8)
    p_sec3.paragraph_format.space_after = Pt(8)
    run_sec3 = p_sec3.add_run("SCREEN SHOT OF OUTPUT")
    run_sec3.font.name = "Times New Roman"
    run_sec3.font.size = Pt(15)
    run_sec3.font.bold = True

    if os.path.exists("screenshot_ml_xai.png"):
        p_img1 = doc.add_paragraph()
        p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img1.paragraph_format.space_after = Pt(2)
        p_img1.add_run().add_picture("screenshot_ml_xai.png", width=Inches(5.8))

        p_cap1 = doc.add_paragraph()
        p_cap1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap1.paragraph_format.space_after = Pt(8)
        run_cap1 = p_cap1.add_run("Figure 1: Multi-Model Machine Learning Ensemble Risk Breakdown & Local Explainable AI (XAI) Attribution Scores")
        run_cap1.font.italic = True
        run_cap1.font.size = Pt(9.5)

    if os.path.exists("screenshot_mri_pipeline.png"):
        p_img2 = doc.add_paragraph()
        p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img2.paragraph_format.space_after = Pt(2)
        p_img2.add_run().add_picture("screenshot_mri_pipeline.png", width=Inches(5.8))

        p_cap2 = doc.add_paragraph()
        p_cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap2.paragraph_format.space_after = Pt(8)
        run_cap2 = p_cap2.add_run("Figure 2: Automated 5-Stage Axial T2-FLAIR MRI Brain Infarct Segmentation & Color Delineation Overlay")
        run_cap2.font.italic = True
        run_cap2.font.size = Pt(9.5)

    if os.path.exists("screenshot_nlp_emergency.png"):
        p_img3 = doc.add_paragraph()
        p_img3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img3.paragraph_format.space_after = Pt(2)
        p_img3.add_run().add_picture("screenshot_nlp_emergency.png", width=Inches(5.8))

        p_cap3 = doc.add_paragraph()
        p_cap3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap3.paragraph_format.space_after = Pt(6)
        run_cap3 = p_cap3.add_run("Figure 3: Clinical Text NLP Entity Extraction with Automated ICD-10/SNOMED Mapping & Emergency Alerting")
        run_cap3.font.italic = True
        run_cap3.font.size = Pt(9.5)

    doc.add_page_break()

    # ---------------- PAGE 7: CONCLUSION ----------------
    p_sec4 = doc.add_paragraph()
    p_sec4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sec4.paragraph_format.space_before = Pt(8)
    p_sec4.paragraph_format.space_after = Pt(10)
    run_sec4 = p_sec4.add_run("CONCLUSION")
    run_sec4.font.name = "Times New Roman"
    run_sec4.font.size = Pt(15)
    run_sec4.font.bold = True

    c_points = [
        ("4.1 Summary of Technical Achievements",
         "The CogniCare AI platform successfully delivers an integrated, multi-modal clinical intelligence environment tailored for brain stroke prevention, "
         "early diagnosis, and emergency orchestration. By harmonizing machine learning predictive classification, computer vision segmentation, "
         "and biomedical natural language processing, the system addresses every pivotal stage along the stroke care continuum. "
         "The soft-voting ensemble model combines Random Forest, Gradient Boosting, and AdaBoost to achieve superior generalization, robustness against class imbalance, "
         "and calibrated probability assessments. Concurrently, the 5-stage MRI segmentation module isolates acute ischemic hyperintensities with deterministic precision, "
         "generating crucial quantitative metrics including lesion burden percentage and ASPECTS territory mappings in sub-second execution times."),

        ("4.2 Clinical Significance & Real-World Utility",
         "Through the integration of Explainable AI (XAI) feature attribution, CogniCare AI provides clinicians with transparent rationale for every assessment, "
         "demystifying machine learning predictions and fostering responsible adoption in high-acuity environments. "
         "Furthermore, the clinical NLP assistant reduces manual documentation friction by standardizing unstructured clinical narratives into structured ICD-10 and SNOMED-CT taxonomies. "
         "Most critically, the built-in FAST screening protocol and automated 4.5-hour golden window countdown timer operationalize emergency protocols, "
         "empowering emergency teams and caregivers to expedite thrombolysis and endovascular transfer."),

        ("4.3 Limitations and Future Scope",
         "While the current platform exhibits high accuracy and responsiveness, future development will expand its diagnostic envelope. "
         "Immediate future work includes transitioning the computer vision pipeline to 3D convolutional architectures (such as 3D U-Net and nnU-Net) "
         "to enable volumetric multi-slice perfusion-diffusion mismatch analysis across multi-parametric MRI (DWI, ADC, T2-FLAIR) and CT Angiography (CTA). "
         "Additionally, integrating secure HL7/FHIR APIs will allow seamless, bi-directional electronic health record synchronization with hospital information systems. "
         "Finally, multi-institutional federated learning will be explored to iteratively train and refine models across distributed clinical cohorts while rigorously preserving patient privacy.")
    ]

    for ch_title, ch_body in c_points:
        p_c = doc.add_paragraph()
        p_c.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_c.paragraph_format.space_after = Pt(6)
        r_ct = p_c.add_run(ch_title + "\n")
        r_ct.font.bold = True
        r_ct.font.size = Pt(11.5)
        r_cb = p_c.add_run(ch_body)
        r_cb.font.size = Pt(11)

    docx_path = "EXP_10_CogniCare_AI_Mini_Project_Report.docx"
    doc.save(docx_path)
    print(f"Successfully generated DOCX report: {docx_path}")

def build_pdf_report():
    pdf_filename = "EXP_10_CogniCare_AI_Mini_Project_Report.pdf"
    
    # 0.5 in top/bottom, 0.7 in left/right
    doc = SimpleDocTemplate(
        pdf_filename,
        pagesize=letter,
        leftMargin=50,
        rightMargin=50,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()

    style_center_bold_14 = ParagraphStyle(
        'CenterBold14',
        parent=styles['Normal'],
        fontName='Times-Bold',
        fontSize=14,
        leading=17,
        alignment=TA_CENTER
    )
    style_center_14 = ParagraphStyle(
        'Center14',
        parent=styles['Normal'],
        fontName='Times-Roman',
        fontSize=13,
        leading=16,
        alignment=TA_CENTER
    )
    style_title = ParagraphStyle(
        'ProjectTitle',
        parent=styles['Normal'],
        fontName='Times-Bold',
        fontSize=15,
        leading=20,
        alignment=TA_CENTER
    )
    style_center_12 = ParagraphStyle(
        'Center12',
        parent=styles['Normal'],
        fontName='Times-Roman',
        fontSize=11.5,
        leading=15,
        alignment=TA_CENTER
    )
    style_center_bold_12 = ParagraphStyle(
        'CenterBold12',
        parent=styles['Normal'],
        fontName='Times-Bold',
        fontSize=11.5,
        leading=15,
        alignment=TA_CENTER
    )
    style_section_h1 = ParagraphStyle(
        'SectionH1',
        parent=styles['Normal'],
        fontName='Times-Bold',
        fontSize=14,
        leading=18,
        alignment=TA_CENTER,
        spaceAfter=10
    )
    style_body_intro = ParagraphStyle(
        'BodyIntro',
        parent=styles['Normal'],
        fontName='Times-Roman',
        fontSize=9.5,
        leading=13.0,
        alignment=TA_JUSTIFY,
        spaceAfter=4
    )
    style_body = ParagraphStyle(
        'BodyJustify',
        parent=styles['Normal'],
        fontName='Times-Roman',
        fontSize=10,
        leading=14.0,
        alignment=TA_JUSTIFY,
        spaceAfter=6
    )
    style_caption = ParagraphStyle(
        'Caption',
        parent=styles['Normal'],
        fontName='Times-Italic',
        fontSize=8.5,
        leading=11,
        alignment=TA_CENTER,
        spaceAfter=6
    )

    story = []

    # ---------------- PAGE 1: TITLE PAGE ----------------
    story.append(Paragraph("EXP 10", style_center_bold_14))
    story.append(Spacer(1, 10))
    story.append(Paragraph("Report of Mini-project On", style_center_14))
    story.append(Spacer(1, 12))
    story.append(Paragraph("CogniCare AI: Brain Stroke Risk Prediction, MRI Lesion Segmentation and Emergency Response Platform", style_title))
    story.append(Spacer(1, 14))
    
    p_subm = (
        "Submitted in partial fulfillment of the requirements of the Mini project in the<br/>"
        "<b>Subject: AI FOR HEALTHCARE of</b><br/>"
        "Semester VII, FINAL Year Artificial Intelligence and Data Science"
    )
    story.append(Paragraph(p_subm, style_center_12))
    story.append(Spacer(1, 12))
    story.append(Paragraph("by", style_center_12))
    story.append(Spacer(1, 2))
    story.append(Paragraph("<b>Shruti Gauchandra (Roll No. 15)</b>", style_center_bold_12))
    story.append(Spacer(1, 12))

    if os.path.exists("page_1_img_1.png"):
        story.append(RLImage("page_1_img_1.png", width=70, height=70))
        story.append(Spacer(1, 6))

    story.append(Paragraph("<b>University of Mumbai</b>", style_center_bold_12))
    story.append(Spacer(1, 3))
    story.append(Paragraph("<b>Vidyavardhini's College of Engineering & Technology</b>", style_center_bold_12))
    story.append(Spacer(1, 3))
    story.append(Paragraph("<b>Department of Artificial Intelligence and Data Science</b>", style_center_bold_12))
    story.append(Spacer(1, 8))

    if os.path.exists("page_1_img_2.png"):
        story.append(RLImage("page_1_img_2.png", width=70, height=72))
        story.append(Spacer(1, 6))

    story.append(Paragraph("<b>(A.Y. 2026-27)</b>", style_center_bold_12))
    story.append(PageBreak())

    # ---------------- PAGE 2: CERTIFICATE ----------------
    story.append(Spacer(1, 24))
    story.append(Paragraph("CERTIFICATE", style_section_h1))
    story.append(Spacer(1, 20))
    
    cert_text = (
        "This is to certify that the Mini Project entitled <b>\"CogniCare AI: Brain Stroke Risk Prediction, "
        "MRI Lesion Segmentation and Emergency Response Platform\"</b> is submitted by "
        "<b>Shruti Gauchandra (Roll No. 15)</b> for the subject of <b>AI for Healthcare</b> in the "
        "<b>Department of Artificial Intelligence and Data Science</b> as a record of work done by him/her "
        "under our supervision and guidance."
    )
    story.append(Paragraph(cert_text, ParagraphStyle('CertBody', parent=style_body, fontSize=11.5, leading=17)))
    story.append(Spacer(1, 50))

    guide_text = (
        "<b>Guide</b><br/>"
        "<b>Asst. Prof. Kranti Gule</b><br/>"
        "Department of Artificial Intelligence & Data Science<br/>"
        "Vidyavardhini's College of Engineering & Technology"
    )
    story.append(Paragraph(guide_text, ParagraphStyle('Guide', parent=styles['Normal'], fontName='Times-Roman', fontSize=11.5, leading=16)))
    story.append(PageBreak())

    # ---------------- PAGE 3: CONTENTS / INDEX ----------------
    story.append(Spacer(1, 20))
    story.append(Paragraph("Contents", style_section_h1))
    story.append(Spacer(1, 15))

    toc_data = [
        [Paragraph("<b>Sr. No</b>", style_center_bold_12), Paragraph("<b>Title</b>", style_center_bold_12), Paragraph("<b>Page No</b>", style_center_bold_12)],
        [Paragraph("1", style_center_12), Paragraph("Introduction of Project", style_body), Paragraph("4", style_center_12)],
        [Paragraph("2", style_center_12), Paragraph("Importance of Project", style_body), Paragraph("5", style_center_12)],
        [Paragraph("3", style_center_12), Paragraph("Screen Shot of Output", style_body), Paragraph("6", style_center_12)],
        [Paragraph("4", style_center_12), Paragraph("Conclusion", style_body), Paragraph("7", style_center_12)]
    ]
    toc_table = RLTable(toc_data, colWidths=[60, 360, 80])
    toc_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#f1f5f9')),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('GRID', (0,0), (-1,-1), 1, colors.HexColor('#cbd5e1')),
        ('TOPPADDING', (0,0), (-1,-1), 8),
        ('BOTTOMPADDING', (0,0), (-1,-1), 8),
    ]))
    story.append(toc_table)
    story.append(PageBreak())

    # ---------------- PAGE 4: INTRODUCTION ----------------
    story.append(Paragraph("INTRODUCTION", style_section_h1))
    intro_p1 = (
        "Stroke is one of the leading global causes of mortality and acquired adult long-term disability, "
        "accounting for approximately 11% of all deaths annually according to the World Health Organization (WHO). "
        "An acute ischemic stroke (AIS) occurs when a thrombus or embolus occludes cerebral blood vessels, depriving downstream "
        "neural tissue of oxygen and glucose. In clinical stroke neurology, the widely recognized paradigm is 'Time is Brain'—for "
        "every minute an ischemic stroke remains untreated, an estimated 1.9 million neurons, 14 billion synapses, and 12 km "
        "of myelinated nerve fibers undergo irreversible necrosis. Despite the efficacy of contemporary reperfusion therapies "
        "such as intravenous tissue plasminogen activator (IV rtPA) and endovascular thrombectomy, systemic bottlenecks in symptom "
        "recognition, clinical risk estimation, radiologic interpretation, and emergency center navigation lead to severe delays "
        "and suboptimal patient outcomes."
    )
    story.append(Paragraph(intro_p1, style_body_intro))

    intro_p2 = (
        "To mitigate these critical challenges, CogniCare AI is developed as an end-to-end, multi-modal artificial intelligence "
        "platform dedicated to comprehensive brain stroke risk management, automated neuroimaging analysis, and rapid emergency dispatch. "
        "Designed in strict accordance with the AI for Healthcare curriculum, the system bridges the gap between predictive statistical modeling, "
        "computer vision, medical natural language processing (NLP), and frontline clinical workflows. "
        "CogniCare AI is architected into six key cohesive modules:"
    )
    story.append(Paragraph(intro_p2, style_body_intro))

    modules_info = [
        ("<b>1. Public Awareness & FAST Screening Engine: </b>"
         "An interactive digital screening assessment implementing the clinical Face, Arm, Speech, and Time (FAST) diagnostic protocol. "
         "The tool captures self-reported or bystander-observed deficits and immediately categorizes the suspicion index, initiating priority emergency pathways when acute focal neurological deficits are present."),
        ("<b>2. Soft-Voting Machine Learning Ensemble: </b>"
         "A tripartite ensemble fusing Random Forest (n=100 estimators), Gradient Boosting (learning rate = 0.08), and AdaBoost classifiers. "
         "Trained on stratified clinical parameters (including age, hypertension, heart disease, average blood glucose level, BMI, work type, and smoking status), "
         "the ensemble computes well-calibrated stroke probability estimates with high sensitivity."),
        ("<b>3. 5-Stage MRI Brain Scan Segmentation Pipeline: </b>"
         "A deterministic computer vision workflow engineered for axial T2-FLAIR and DWI brain MRI scans. The pipeline sequentially applies: "
         "(i) Grayscale intensity transformation, (ii) Gaussian spatial smoothing (5x5 kernel, sigma=1.2) to attenuate acquisition noise, "
         "(iii) Otsu's optimal bimodal thresholding to isolate hyperintense ischemic territories, (iv) Mathematical morphological opening and closing "
         "to eliminate skull-bone artifacts and ventricular noise, and (v) Telemetric quantification computing segmented lesion area (pixels), "
         "lesion burden ratio (%), and automated ASPECTS anatomical territory localization."),
        ("<b>4. Clinical Natural Language Processing (NLP) Assistant: </b>"
         "A medical text parsing engine that processes unstructured doctor, triage, and nurse EHR progress notes. "
         "Leveraging regex tokenization, clinical synonym dictionaries, and fuzzy string matching, the module extracts symptoms, diagnoses, "
         "and pharmacological prescriptions, automatically mapping them to standardized ontological codes (ICD-10, RxNorm, and SNOMED-CT)."),
        ("<b>5. Explainable AI (XAI) Biomarker Attribution: </b>"
         "Provides transparent, SHAP-inspired local feature contribution scores for each prediction, highlighting critical clinical biomarker alerts "
         "(e.g., severe hyperglycemia >200 mg/dL, stage-2 hypertension) to foster clinical trust and physician-in-the-loop validation."),
        ("<b>6. Stroke Emergency Dispatch & Telemetry: </b>"
         "A one-tap emergency protocol integrating a 4.5-hour intravenous thrombolysis countdown timer, geolocation-based nearest stroke-certified "
         "comprehensive hospital navigation, and paramedic handoff cards."
        )
    ]
    for m in modules_info:
        story.append(Paragraph(m, style_body_intro))

    story.append(PageBreak())

    # ---------------- PAGE 5: IMPORTANCE OF PROJECT ----------------
    story.append(Paragraph("IMPORTANCE OF PROJECT", style_section_h1))

    imp_points = [
        ("<b>1. Critical Golden Hour Optimization and Door-to-Needle Reduction</b><br/>"
         "In acute ischemic stroke, therapeutic efficacy is governed by the therapeutic golden window: intravenous thrombolysis (IV rtPA) must "
         "be administered within 4.5 hours from the last known well time, and endovascular thrombectomy is indicated within 6 to 24 hours. "
         "Every 15-minute reduction in door-to-needle time results in a 5% relative reduction in 90-day mortality. CogniCare AI directly "
         "compresses this critical timeline through instantaneous FAST identification, immediate emergency dispatch, and live countdown telemetry."),

        ("<b>2. Objective, Deterministic, and Quantitative Neuroimaging Assessment</b><br/>"
         "Manual radiologic interpretation of acute brain scans is prone to inter-observer variability, especially in primary care settings without "
         "24/7 on-call neuroradiologists. The 5-stage automated segmentation pipeline delivers standardized, objective metrics of ischemic core burden "
         "and territory localization within milliseconds, preventing diagnostic hesitation and supporting emergency triage decisions."),

        ("<b>3. Eliminating Black-Box Hesitation via Explainable AI (XAI)</b><br/>"
         "A primary barrier to the clinical adoption of machine learning in healthcare is the opaque nature of complex models. "
         "Clinicians cannot ethically or legally act on unverified predictions. CogniCare AI addresses this challenge by providing granular feature "
         "attributions and transparent voting consensus across three distinct ensemble classifiers, demonstrating clearly whether risk is driven "
         "by vascular hypertension, diabetic microangiopathy, or advancing age."),

        ("<b>4. Streamlining Clinical Workflows & EHR Interoperability</b><br/>"
         "Emergency physicians face severe documentation burden, spending up to two hours on administrative electronic health record (EHR) entries "
         "for every hour of direct patient care. CogniCare AI's clinical NLP assistant instantly normalizes free-text clinical notes into internationally "
         "recognized ICD-10 and SNOMED-CT taxonomies, accelerating patient handoffs between emergency departments and neuro-intensive care units."),

        ("<b>5. Democratizing Advanced Stroke Care in Rural and Underserved Facilities</b><br/>"
         "Tertiary stroke centers are disproportionately clustered in major metropolitan hubs. Rural and community district hospitals often lack "
         "sub-specialty stroke teams. CogniCare AI acts as a sophisticated clinical force multiplier, equipping medical officers with specialist-level "
         "decision support and rapid transfer coordination protocols to ensure equitable healthcare delivery.")
    ]
    for imp in imp_points:
        story.append(Paragraph(imp, style_body))

    story.append(PageBreak())

    # ---------------- PAGE 6: SCREEN SHOT OF OUTPUT ----------------
    story.append(Paragraph("SCREEN SHOT OF OUTPUT", style_section_h1))

    if os.path.exists("screenshot_ml_xai.png"):
        story.append(RLImage("screenshot_ml_xai.png", width=460, height=140))
        story.append(Paragraph("Figure 1: Multi-Model Machine Learning Ensemble Risk Breakdown & Local Explainable AI (XAI) Attribution Scores", style_caption))

    if os.path.exists("screenshot_mri_pipeline.png"):
        story.append(RLImage("screenshot_mri_pipeline.png", width=460, height=120))
        story.append(Paragraph("Figure 2: Automated 5-Stage Axial T2-FLAIR MRI Brain Infarct Segmentation & Color Delineation Overlay", style_caption))

    if os.path.exists("screenshot_nlp_emergency.png"):
        story.append(RLImage("screenshot_nlp_emergency.png", width=460, height=110))
        story.append(Paragraph("Figure 3: Clinical Text NLP Entity Extraction with Automated ICD-10/SNOMED Mapping & Emergency Alerting", style_caption))

    story.append(PageBreak())

    # ---------------- PAGE 7: CONCLUSION ----------------
    story.append(Paragraph("CONCLUSION", style_section_h1))

    c_points = [
        ("<b>4.1 Summary of Technical Achievements</b><br/>"
         "The CogniCare AI platform successfully delivers an integrated, multi-modal clinical intelligence environment tailored for brain stroke prevention, "
         "early diagnosis, and emergency orchestration. By harmonizing machine learning predictive classification, computer vision segmentation, "
         "and biomedical natural language processing, the system addresses every pivotal stage along the stroke care continuum. "
         "The soft-voting ensemble model combines Random Forest, Gradient Boosting, and AdaBoost to achieve superior generalization, robustness against class imbalance, "
         "and calibrated probability assessments. Concurrently, the 5-stage MRI segmentation module isolates acute ischemic hyperintensities with deterministic precision, "
         "generating crucial quantitative metrics including lesion burden percentage and ASPECTS territory mappings in sub-second execution times."),

        ("<b>4.2 Clinical Significance & Real-World Utility</b><br/>"
         "Through the integration of Explainable AI (XAI) feature attribution, CogniCare AI provides clinicians with transparent rationale for every assessment, "
         "demystifying machine learning predictions and fostering responsible adoption in high-acuity environments. "
         "Furthermore, the clinical NLP assistant reduces manual documentation friction by standardizing unstructured clinical narratives into structured ICD-10 and SNOMED-CT taxonomies. "
         "Most critically, the built-in FAST screening protocol and automated 4.5-hour golden window countdown timer operationalize emergency protocols, "
         "empowering emergency teams and caregivers to expedite thrombolysis and endovascular transfer."),

        ("<b>4.3 Limitations and Future Scope</b><br/>"
         "While the current platform exhibits high accuracy and responsiveness, future development will expand its diagnostic envelope. "
         "Immediate future work includes transitioning the computer vision pipeline to 3D convolutional architectures (such as 3D U-Net and nnU-Net) "
         "to enable volumetric multi-slice perfusion-diffusion mismatch analysis across multi-parametric MRI (DWI, ADC, T2-FLAIR) and CT Angiography (CTA). "
         "Additionally, integrating secure HL7/FHIR APIs will allow seamless, bi-directional electronic health record synchronization with hospital information systems. "
         "Finally, multi-institutional federated learning will be explored to iteratively train and refine models across distributed clinical cohorts while rigorously preserving patient privacy.")
    ]

    for c in c_points:
        story.append(Paragraph(c, style_body))

    doc.build(story)
    print(f"Successfully generated PDF report: {pdf_filename}")

if __name__ == '__main__':
    build_docx_report()
    build_pdf_report()
