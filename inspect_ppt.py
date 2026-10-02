import pptx

prs = pptx.Presentation(r'C:\Users\Admin\.gemini\antigravity\brain\870bcad5-0f86-48a7-8a75-a343e86ebacb\.user_uploaded\media_1790815052736.pptx')
for i, slide in enumerate(prs.slides):
    print(f'=== SLIDE {i+1} ===')
    for j, shape in enumerate(slide.shapes):
        txt = shape.text.strip().replace('\n', ' ') if shape.has_text_frame else ''
        print(f'  Shape {j} ({shape.name}, type={shape.shape_type}): "{txt[:80]}"')
