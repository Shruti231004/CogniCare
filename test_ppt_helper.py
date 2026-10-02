import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

TEMPLATE_PATH = r"C:\Users\Admin\.gemini\antigravity\brain\870bcad5-0f86-48a7-8a75-a343e86ebacb\.user_uploaded\media_1790815052736.pptx"
OUTPUT_PATH = r"c:\Shruti\AIH Project\CogniCare_AI_Mini_Project_Presentation.pptx"

prs = pptx.Presentation(TEMPLATE_PATH)

def format_bullet_point(tf, bold_prefix, text, font_size=15, is_first=False, level=0):
    if is_first:
        p = tf.paragraphs[0]
    else:
        p = tf.add_paragraph()
    p.level = level
    p.space_after = Pt(8)
    p.space_before = Pt(2)
    
    if bold_prefix:
        r_bold = p.add_run()
        r_bold.text = bold_prefix + ": " if not bold_prefix.endswith(":") else bold_prefix + " "
        r_bold.font.name = "Times New Roman"
        r_bold.font.size = Pt(font_size)
        r_bold.font.bold = True
        r_bold.font.color.rgb = RGBColor(0x0f, 0x17, 0x2a)
        
    r_text = p.add_run()
    r_text.text = text
    r_text.font.name = "Times New Roman"
    r_text.font.size = Pt(font_size)
    r_text.font.bold = False
    r_text.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

print("Helper loaded successfully")
